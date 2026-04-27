"""Load the company catalog from the processed pipeline Excel."""
from __future__ import annotations
import os
import re
from dataclasses import dataclass
from typing import Iterable
import openpyxl

EXCEL_PATH_DEFAULT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "pipeline_output",
    "companies_pipeline.xlsx",
)

# SIC code -> archetype. Only "shift" exists today; the rest are placeholders
# for future archetype modules and currently fall back to "shift".
SIC_ARCHETYPE = {
    "78200": "shift",  # Temp employment agency
    "88100": "shift",  # Social work without accommodation for the elderly/disabled
    "87100": "shift",  # Residential nursing care
    "53202": "shift",  # Courier (will become "mileage" when that archetype lands)
    "49410": "shift",  # Freight road transport
    "98000": "shift",  # Residents property mgmt (will become "twa" when that lands)
}
DEFAULT_ARCHETYPE = "shift"


@dataclass(frozen=True)
class Company:
    company_number: str
    company_name: str
    sic_codes: str
    domain: str
    support_email: str
    archetype: str
    short_name: str = ""          # ← NEW: curated app name from Excel

    @property
    def flavor(self) -> str:
        """Gradle flavor name — must match [a-zA-Z][a-zA-Z0-9]*."""
        return f"c{self.company_number}"

    @property
    def application_id(self) -> str:
        return f"uk.c{self.company_number}.shift"

    @property
    def display_name(self) -> str:
        """Launcher label — prefers curated short_name; falls back to
        title-cased company name with LTD/LIMITED stripped, 30-char cap."""
        # ── Option 1: use the curated short name if present ──
        if self.short_name:
            return self.short_name[:30].rstrip()

        # ── Fallback: auto-generate from company_name ──
        name = self.company_name.strip()
        name = re.sub(r"\s+(LTD|LIMITED|LLP|PLC)\.?$", "", name, flags=re.IGNORECASE)
        name = name.title()
        return name[:30].rstrip()


def load_companies(
    excel_path: str = EXCEL_PATH_DEFAULT,
    limit: int | None = None,
) -> list[Company]:
    """Read companies from the pipeline Excel. Skips rows without a company number."""
    wb = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)

    def col(name: str) -> int:
        for i, h in enumerate(header):
            if h and h.strip().lower() == name.lower():
                return i
        raise KeyError(f"column not found: {name}")

    def col_optional(name: str) -> int | None:
        """Return column index or None if the column doesn't exist yet."""
        for i, h in enumerate(header):
            if h and h.strip().lower() == name.lower():
                return i
        return None

    idx_num   = col("Company Number")
    idx_name  = col("Company Name")
    idx_sic   = col("SIC Codes")
    idx_domain = col("Domain")
    idx_email = col("Assigned Email")
    idx_short = col_optional("Short Name")  # ← NEW: optional column

    out: list[Company] = []
    for row in rows:
        if not row or not row[idx_num]:
            continue

        cn = str(row[idx_num]).strip()
        # Pad to 8 digits (Companies House format)
        if cn.isdigit() and len(cn) < 8:
            cn = cn.zfill(8)

        sic_raw = (row[idx_sic] or "")
        primary_sic = sic_raw.split(",")[0].strip() if sic_raw else ""
        archetype = SIC_ARCHETYPE.get(primary_sic, DEFAULT_ARCHETYPE)

        domain = (row[idx_domain] or "").strip() if row[idx_domain] else ""
        email  = (row[idx_email] or "").strip() if row[idx_email] else ""
        support = f"support@{domain}" if domain else (email or "support@example.uk")

        # ── Read the optional Short Name cell ──
        short = ""
        if idx_short is not None:
            short = (row[idx_short] or "").strip() if row[idx_short] else ""

        out.append(Company(
            company_number=cn,
            company_name=(row[idx_name] or "").strip(),
            sic_codes=sic_raw,
            domain=domain,
            support_email=support,
            archetype=archetype,
            short_name=short,                        # ← NEW
        ))

        if limit is not None and len(out) >= limit:
            break

    wb.close()
    return out


def summarize(companies: Iterable[Company]) -> str:
    lines = [f"{'#':>3} {'Flavor':<12} {'SIC':<6} {'Arch':<6} Name"]
    for i, c in enumerate(companies, 1):
        primary_sic = c.sic_codes.split(",")[0].strip() if c.sic_codes else "-"
        lines.append(f"{i:>3} {c.flavor:<12} {primary_sic:<6} {c.archetype:<6} {c.display_name}")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys
    n = int(sys.argv[1]) if len(sys.argv) > 1 else 3
    cs = load_companies(limit=n)
    print(summarize(cs))
